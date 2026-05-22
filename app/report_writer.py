from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

HF = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
FF = Font(color='FFFFFF', bold=True)


def write_report(out: Path, payload: dict, meta: dict, process_log: list[str]):
    with pd.ExcelWriter(out, engine='openpyxl') as wr:
        resumen = pd.DataFrame([
            ['Fecha y hora', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
            ['Archivo1', meta['archivo1']], ['Archivo2', meta['archivo2']],
            ['Total registros Archivo1', meta['total_fallecidos']], ['Total registros Archivo2', meta['total_citas']],
            ['Total citas analizadas', meta['total_citas_analizadas']], ['Filtro solo pendientes', meta['solo_pendientes']],
            ['Total coincidencias cédula', len(payload['cedula'])], ['Total coincidencias nombre exacto', len(payload['nombre_exacto'])],
            ['Total posibles coincidencias nombre', len(payload['posible_nombre'])], ['Total cédulas repetidas Archivo1', meta['dup_cedula_a1']],
            ['Total cédulas repetidas Archivo2', meta['dup_cedula_a2']], ['Total registros cédula inválida', meta['cedulas_invalidas']],
            ['Total casos revisión administrativa', meta['revision']],
        ], columns=['Métrica', 'Valor'])
        resumen.to_excel(wr, sheet_name='RESUMEN', index=False)
        payload['cedula'].to_excel(wr, sheet_name='COINCIDENCIAS_CEDULA_100', index=False)
        payload['cedula_nombre_diff'].to_excel(wr, sheet_name='CEDULA_COINCIDE_NOMBRE_DIFEREN', index=False)
        payload['nombre_exacto'].to_excel(wr, sheet_name='NOMBRE_100_VERIFICAR', index=False)
        payload['posible_nombre'].to_excel(wr, sheet_name='POSIBLES_COINCIDENCIAS_NOMBRE', index=False)
        _dups(payload['fallecidos'], 'cedula_norm', 'nom_norm', 'Archivo1').to_excel(wr, sheet_name='DUPLICADOS_ARCHIVO1', index=False)
        _dups(payload['citas'], 'cedula_norm', 'nom_norm', 'Archivo2').to_excel(wr, sheet_name='DUPLICADOS_ARCHIVO2', index=False)
        payload['sin_coincidencia'].to_excel(wr, sheet_name='SIN_COINCIDENCIA', index=False)
        pd.DataFrame({'LOG': process_log}).to_excel(wr, sheet_name='LOG_PROCESO', index=False)

        for ws in wr.book.worksheets:
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = 'A2'
            for c in ws[1]:
                c.fill = HF; c.font = FF
            for col in ws.columns:
                w = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(w+2, 50)


def _dups(df, ccol, ncol, tag):
    a = df[df[ccol].astype(str).str.len()>0].groupby(ccol).size().reset_index(name='conteo')
    a = a[a['conteo']>1]; a['tipo']=f'Cedula repetida {tag}'
    b = df[df[ncol].astype(str).str.len()>0].groupby(ncol).size().reset_index(name='conteo')
    b = b[b['conteo']>1]; b['tipo']=f'Nombre repetido {tag}'; b = b.rename(columns={ncol:ccol})
    return pd.concat([a[[ccol,'conteo','tipo']], b[[ccol,'conteo','tipo']]], ignore_index=True)
